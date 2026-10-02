from __future__ import annotations

import datetime
import uuid
from typing import List, Optional

from fastapi import HTTPException, status
from sqlalchemy.orm import Session

from app.models.debt import SoCongNo
from app.models.contract import HopDong, Phi
from app.models.dorm import Giuong, Phong, Tang, ToaNha
from app.models.invoice import HoaDon, LoaiHoaDon, TrangThaiHoaDon
from app.models.user import NguoiDung, SinhVien
from app.schemas.invoice import (
    HoaDonResponse,
    PublishResultResponse,
    RoomBillingCandidate,
    RoomInvoicePublishRequest,
    UtilityBillingCandidate,
    UtilityInvoicePublishRequest,
    UtilityMeterUploadResponse,
)

DEFAULT_DEMO_UTILITY_ROOMS = [
    {
        "ma_phong": "P102",
        "so_phong": "102",
        "toa_nha": "A1",
        "so_sinh_vien": 4,
        "chi_so_dien_cu_moi": "1240 - 1340",
        "so_dien_kwh": 100,
        "chi_so_nuoc_cu_moi": "450 - 460",
        "so_nuoc_m3": 10,
        "tong_tien": 450000.0,
    },
    {
        "ma_phong": "P103",
        "so_phong": "103",
        "toa_nha": "A1",
        "so_sinh_vien": 4,
        "chi_so_dien_cu_moi": "2100 - 2215",
        "so_dien_kwh": 115,
        "chi_so_nuoc_cu_moi": "610 - 622",
        "so_nuoc_m3": 12,
        "tong_tien": 525000.0,
    },
    {
        "ma_phong": "P201",
        "so_phong": "201",
        "toa_nha": "A2",
        "so_sinh_vien": 4,
        "chi_so_dien_cu_moi": "0890 - 0985",
        "so_dien_kwh": 95,
        "chi_so_nuoc_cu_moi": "320 - 328",
        "so_nuoc_m3": 8,
        "tong_tien": 405000.0,
    },
    {
        "ma_phong": "P205",
        "so_phong": "205",
        "toa_nha": "A2",
        "so_sinh_vien": 4,
        "chi_so_dien_cu_moi": "1540 - 1670",
        "so_dien_kwh": 130,
        "chi_so_nuoc_cu_moi": "540 - 554",
        "so_nuoc_m3": 14,
        "tong_tien": 600000.0,
    },
]


class InvoiceService:
    @staticmethod
    def get_room_billing_candidates(
        db: Session,
        ky_thanh_toan: str = "Tháng 09/2026",
        don_gia_thang: float = 350000.0,
        thoi_gian_o_thang: int = 1,
        ap_dung: Optional[str] = None,
    ) -> List[RoomBillingCandidate]:
        """
        Lấy danh sách sinh viên nội trú cần lập hóa đơn tiền phòng theo tháng:
        - Đơn giá niêm yết: Phòng tiêu chuẩn là 350.000 VNĐ/tháng, Phòng dịch vụ là 650.000 VNĐ/tháng.
        - Phạm vi áp dụng: Tất cả phòng, Chỉ phòng tiêu chuẩn, Chỉ phòng dịch vụ (hoặc theo tòa).
        - Chỉ lấy từ các hợp đồng ACTIVE thực sự tồn tại trong CSDL.
        """
        # 1. Truy vấn các hợp đồng đang hiệu lực từ CSDL
        query = db.query(HopDong).filter(HopDong.trang_thai == "ACTIVE")
        active_contracts = query.all()

        candidates: List[RoomBillingCandidate] = []

        for hd in active_contracts:
            sv = hd.sinh_vien
            nd = sv.nguoi_dung if sv else None
            ho_ten = nd.ho_ten if nd else (sv.msv if sv else hd.msv)

            # Tìm thông tin phòng & loại phòng
            phong_str = "Chưa xếp phòng"
            ma_phong_val = None
            loai_phong_label = "Phòng tiêu chuẩn"
            unit_price = 350000.0

            if hd.giuong and hd.giuong.phong:
                p = hd.giuong.phong
                ma_phong_val = p.ma_phong
                t = p.tang.toa_nha.ten_toa if p.tang and p.tang.toa_nha else ""
                phong_str = f"P{p.so_phong} - {t}".strip(" -")
                raw_loai = (p.loai_phong or "").lower()
                if "dịch vụ" in raw_loai or "dich vu" in raw_loai or "service" in raw_loai:
                    loai_phong_label = "Phòng dịch vụ"
                    unit_price = 650000.0
                else:
                    loai_phong_label = "Phòng tiêu chuẩn"
                    unit_price = 350000.0

            # Lọc theo phạm vi áp dụng (Chỉ phòng tiêu chuẩn / Chỉ phòng dịch vụ)
            if ap_dung:
                ap_norm = ap_dung.lower()
                if "tiêu chuẩn" in ap_norm and "dịch vụ" not in ap_norm:
                    if loai_phong_label != "Phòng tiêu chuẩn":
                        continue
                elif "dịch vụ" in ap_norm and "tiêu chuẩn" not in ap_norm:
                    if loai_phong_label != "Phòng dịch vụ":
                        continue
                elif "tòa" in ap_norm or "toa" in ap_norm:
                    import re
                    m = re.search(r"A\d+", ap_dung, re.IGNORECASE)
                    if m and m.group(0).upper() not in phong_str.upper():
                        continue

            # Định dạng thời hạn hợp đồng
            start_str = hd.ngay_bat_dau.strftime("%d/%m/%Y") if hd.ngay_bat_dau else ""
            end_str = hd.ngay_ket_thuc.strftime("%d/%m/%Y") if hd.ngay_ket_thuc else "Hiện tại"
            term_str = f"{start_str} - {end_str}".strip(" -")

            # Kiểm tra xem hợp đồng này đã lập hóa đơn chưa cho kỳ này
            existing_inv = (
                db.query(HoaDon)
                .filter(
                    HoaDon.ma_hop_dong == hd.ma_hop_dong,
                    HoaDon.loai_hoa_don == LoaiHoaDon.TIEN_PHONG.value,
                    HoaDon.ky_thanh_toan == ky_thanh_toan,
                )
                .first()
            )

            current_amount = unit_price * thoi_gian_o_thang

            candidates.append(
                RoomBillingCandidate(
                    msv=hd.msv,
                    ho_ten=ho_ten,
                    phong=phong_str,
                    loai_phong=loai_phong_label,
                    ma_phong=ma_phong_val,
                    ma_hop_dong=hd.ma_hop_dong,
                    thoi_han_hop_dong=term_str,
                    thoi_gian_o_thang=thoi_gian_o_thang,
                    don_gia_thang=unit_price,
                    so_tien=current_amount,
                    da_lap_hoa_don=existing_inv is not None,
                )
            )

        return candidates

    @staticmethod
    def publish_room_invoices(
        db: Session,
        request: RoomInvoicePublishRequest,
        creator_username: str,
    ) -> PublishResultResponse:
        """
        Phát hành đợt hóa đơn tiền phòng định kỳ theo kỳ học.
        Chỉ phát hành hóa đơn với các hợp đồng thực sự tồn tại trong CSDL.
        Kiểm tra chặt chẽ tính hợp lệ, toàn vẹn khóa ngoại (FK) và chống tạo trùng lặp hóa đơn.
        """
        if not request.ky_thanh_toan or not request.ky_thanh_toan.strip():
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Kỳ thanh toán không được để trống",
            )

        if not request.han_thanh_toan:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Hạn chót thanh toán không được để trống",
            )

        # 1. Truy vấn các hợp đồng ACTIVE thực tế trong CSDL
        query = db.query(HopDong).filter(HopDong.trang_thai == "ACTIVE")
        if request.danh_sach_msv:
            query = query.filter(HopDong.msv.in_(request.danh_sach_msv))
        valid_contracts = query.all()

        # Lọc theo kỳ thanh toán nếu không truyền danh sách MSV cụ thể
        if not request.danh_sach_msv:
            import re
            m_month = re.search(r"(\d{1,2})/(\d{4})", request.ky_thanh_toan)
            if m_month:
                month_num = int(m_month.group(1))
                year_num = int(m_month.group(2))
                import calendar
                _, last_day = calendar.monthrange(year_num, month_num)
                period_start = datetime.date(year_num, month_num, 1)
                period_end = datetime.date(year_num, month_num, last_day)
                valid_contracts = [
                    hd for hd in valid_contracts
                    if (hd.ngay_bat_dau or datetime.date(2000, 1, 1)) <= period_end
                    and (hd.ngay_ket_thuc or datetime.date(2099, 12, 31)) >= period_start
                ]
            else:
                years = [int(y) for y in re.findall(r"\b(20\d\d)\b", request.ky_thanh_toan)]
                if years:
                    start_year = min(years)
                    end_year = max(years)
                    period_start = datetime.date(start_year, 8, 1)
                    period_end = datetime.date(end_year, 7, 31) if end_year > start_year else datetime.date(start_year, 12, 31)
                    valid_contracts = [
                        hd for hd in valid_contracts
                        if (hd.ngay_bat_dau or datetime.date(2000, 1, 1)) <= period_end
                        and (hd.ngay_ket_thuc or datetime.date(2099, 12, 31)) >= period_start
                    ]

        # Lọc theo tòa nhà áp dụng nếu không truyền danh sách MSV cụ thể
        if not request.danh_sach_msv and request.ap_dung and "tòa" in request.ap_dung.lower():
            import re
            m = re.search(r"A\d+", request.ap_dung, re.IGNORECASE)
            if m:
                target_building = m.group(0).upper()
                valid_contracts = [
                    hd for hd in valid_contracts
                    if hd.giuong and hd.giuong.phong and hd.giuong.phong.tang and hd.giuong.phong.tang.toa_nha
                    and target_building in f"{hd.giuong.phong.tang.toa_nha.ma_toa} {hd.giuong.phong.tang.toa_nha.ten_toa}".upper()
                ]

        # Xác thực tính tồn tại của hợp đồng
        if request.danh_sach_msv:
            found_msvs = {c.msv for c in valid_contracts}
            missing_msvs = [msv for msv in request.danh_sach_msv if msv not in found_msvs]
            if missing_msvs:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"Không tìm thấy hợp đồng hợp lệ hoặc hợp đồng không tồn tại trong hệ thống cho sinh viên: {', '.join(missing_msvs)}. Chỉ phát hành hóa đơn với hợp đồng tồn tại.",
                )

        if not valid_contracts:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Không tìm thấy hợp đồng hợp lệ nào tồn tại trong hệ thống để phát hành hóa đơn. Chỉ phát hành hóa đơn với hợp đồng tồn tại.",
            )

        created_invoices: List[HoaDon] = []
        already_issued_info: List[str] = []

        try:
            for hd in valid_contracts:
                # Trích xuất thông tin người dùng và phòng từ quan hệ hợp đồng
                sv = hd.sinh_vien
                nd = sv.nguoi_dung if sv else None
                ho_ten = nd.ho_ten if nd else (sv.msv if sv else hd.msv)

                phong_str = "Chưa xếp phòng"
                ma_phong_val = None
                loai_phong_label = "Phòng tiêu chuẩn"
                unit_price = 350000.0

                if hd.giuong and hd.giuong.phong:
                    p = hd.giuong.phong
                    ma_phong_val = p.ma_phong
                    t = p.tang.toa_nha.ten_toa if p.tang and p.tang.toa_nha else ""
                    phong_str = f"P{p.so_phong} - {t}".strip(" -")
                    raw_loai = (p.loai_phong or "").lower()
                    if "dịch vụ" in raw_loai or "dich vu" in raw_loai or "service" in raw_loai:
                        loai_phong_label = "Phòng dịch vụ"
                        unit_price = 650000.0
                    else:
                        loai_phong_label = "Phòng tiêu chuẩn"
                        unit_price = 350000.0

                # Lọc theo phạm vi áp dụng (Chỉ phòng tiêu chuẩn / Chỉ phòng dịch vụ)
                if request.ap_dung:
                    ap_norm = request.ap_dung.lower()
                    if "tiêu chuẩn" in ap_norm and "dịch vụ" not in ap_norm and loai_phong_label != "Phòng tiêu chuẩn":
                        continue
                    if "dịch vụ" in ap_norm and "tiêu chuẩn" not in ap_norm and loai_phong_label != "Phòng dịch vụ":
                        continue

                # 2. Kiểm tra chống tạo trùng lặp theo hợp đồng và kỳ thanh toán
                existing = (
                    db.query(HoaDon)
                    .filter(
                        HoaDon.ma_hop_dong == hd.ma_hop_dong,
                        HoaDon.loai_hoa_don == LoaiHoaDon.TIEN_PHONG.value,
                        HoaDon.ky_thanh_toan == request.ky_thanh_toan,
                    )
                    .first()
                )

                if existing:
                    already_issued_info.append(f"{hd.msv} (HĐ: {hd.ma_hop_dong})")
                    continue

                current_student_amount = unit_price * request.thoi_gian_o_thang

                # Sinh mã hóa đơn duy nhất dạng HDTP-YYYYMMDD-UUID4[:6]
                today_str = datetime.date.today().strftime("%Y%m%d")
                unique_suffix = uuid.uuid4().hex[:6].upper()
                ma_hd = f"HDTP-{today_str}-{unique_suffix}"

                # 3. Tạo hóa đơn tiền phòng với ma_hop_dong thực tế từ bảng hop_dong (tuân thủ FK 100%)
                new_invoice = HoaDon(
                    ma_hoa_don=ma_hd,
                    ma_hop_dong=hd.ma_hop_dong,
                    msv=hd.msv,
                    ho_ten=ho_ten,
                    ma_phong=ma_phong_val,
                    so_phong=phong_str,
                    loai_hoa_don=LoaiHoaDon.TIEN_PHONG.value,
                    ky_thanh_toan=request.ky_thanh_toan,
                    so_tien=current_student_amount,
                    ngay_lap=datetime.date.today(),
                    han_thanh_toan=request.han_thanh_toan,
                    trang_thai=TrangThaiHoaDon.CHUA_THANH_TOAN.value,
                    ghi_chu=request.ghi_chu or f"Hóa đơn tiền phòng {request.ky_thanh_toan} ({loai_phong_label})",
                    nguoi_tao=creator_username,
                    ngay_tao=datetime.datetime.now(),
                )
                db.add(new_invoice)

                # 4. Tạo bản ghi vào bảng phi hiện có gắn với mã hợp đồng thực tế
                phi_rec = Phi(
                    ma_phi=f"P-{unique_suffix}"[:20],
                    ma_hop_dong=hd.ma_hop_dong,
                    loai_phi="TIEN_PHONG",
                    so_tien=current_student_amount,
                    han_nop=request.han_thanh_toan,
                )
                db.add(phi_rec)

                # 5. Tạo bản ghi vào bảng so_cong_no tương ứng
                so_cong_no_rec = SoCongNo(
                    ma_hoa_don=new_invoice.ma_hoa_don,
                    msv=hd.msv,
                    so_phong=phong_str,
                    loai_cong_no=new_invoice.loai_hoa_don,
                    ky_thanh_toan=new_invoice.ky_thanh_toan,
                    tong_tien=new_invoice.so_tien,
                    da_tra=0.0,
                    con_thieu=new_invoice.so_tien,
                    han_thanh_toan=new_invoice.han_thanh_toan,
                    trang_thai="CON_NO",
                    ghi_chu=new_invoice.ghi_chu or f"Công nợ tiền phòng {new_invoice.ky_thanh_toan}",
                    ngay_cap_nhat=datetime.datetime.now(),
                )
                db.add(so_cong_no_rec)

                created_invoices.append(new_invoice)

            if not created_invoices and already_issued_info:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"Tất cả hợp đồng được chọn đã được phát hành hóa đơn cho {request.ky_thanh_toan} trước đó ({', '.join(already_issued_info)}). Không thể phát hành trùng!",
                )

            db.commit()

            # Refresh các bản ghi để lấy dữ liệu hoàn chỉnh
            for inv in created_invoices:
                db.refresh(inv)
        except HTTPException:
            raise
        except Exception as e:
            db.rollback()
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail=f"Lỗi khi phát hành hóa đơn: {str(e)}",
            )

        invoices_response = [
            HoaDonResponse(
                ma_hoa_don=inv.ma_hoa_don,
                ma_hop_dong=inv.ma_hop_dong,
                msv=inv.msv,
                ho_ten=inv.ho_ten,
                ma_phong=inv.ma_phong,
                so_phong=inv.so_phong,
                loai_hoa_don=inv.loai_hoa_don,
                ky_thanh_toan=inv.ky_thanh_toan,
                so_tien=inv.so_tien,
                ngay_lap=inv.ngay_lap,
                han_thanh_toan=inv.han_thanh_toan,
                trang_thai=inv.trang_thai,
                ghi_chu=inv.ghi_chu,
                nguoi_tao=inv.nguoi_tao,
            )
            for inv in created_invoices
        ]

        total_sum = sum(inv.so_tien for inv in created_invoices)
        msg = f"Đã phát hành thành công {len(created_invoices)} hóa đơn tiền phòng cho {request.ky_thanh_toan}."
        if already_issued_info:
            msg += f" (Đã bỏ qua {len(already_issued_info)} hợp đồng đã có hóa đơn trước đó: {', '.join(already_issued_info)})"

        return PublishResultResponse(
            success=True,
            message=msg,
            tong_hoa_don=len(created_invoices),
            tong_so_tien=total_sum,
            invoices=invoices_response,
        )

    @staticmethod
    def get_utility_billing_candidates(
        db: Session,
        thang: str = "Tháng 09/2026",
    ) -> List[UtilityBillingCandidate]:
        """
        Lấy danh sách các phòng phục vụ lập hóa đơn tiền điện nước theo tháng.
        Kiểm tra hóa đơn điện nước đã phát hành cho phòng trong tháng tương ứng.
        """
        # Truy vấn phòng từ CSDL
        rooms = db.query(Phong).all()
        candidates: List[UtilityBillingCandidate] = []

        if rooms:
            for r in rooms:
                toa_ten = r.tang.toa_nha.ten_toa if r.tang and r.tang.toa_nha else "KTX"
                # Đếm số sinh viên đang ở trong phòng
                student_count = (
                    db.query(HopDong)
                    .join(Giuong, HopDong.ma_giuong == Giuong.ma_giuong)
                    .filter(Giuong.ma_phong == r.ma_phong, HopDong.trang_thai == "ACTIVE")
                    .count()
                )

                existing = (
                    db.query(HoaDon)
                    .filter(
                        HoaDon.ma_phong == r.ma_phong,
                        HoaDon.loai_hoa_don == LoaiHoaDon.DIEN_NUOC.value,
                        HoaDon.ky_thanh_toan == thang,
                    )
                    .first()
                )

                candidates.append(
                    UtilityBillingCandidate(
                        ma_phong=r.ma_phong,
                        so_phong=r.so_phong,
                        toa_nha=toa_ten,
                        so_sinh_vien=student_count,
                        chi_so_dien_cu_moi="1000 - 1100",
                        so_dien_kwh=100,
                        chi_so_nuoc_cu_moi="400 - 410",
                        so_nuoc_m3=10,
                        tong_tien=450000.0,
                        da_lap_hoa_don=existing is not None,
                    )
                )
        else:
            # Fallback danh sách phòng mẫu chuẩn Figma
            for item in DEFAULT_DEMO_UTILITY_ROOMS:
                existing = (
                    db.query(HoaDon)
                    .filter(
                        HoaDon.ma_phong == item["ma_phong"],
                        HoaDon.loai_hoa_don == LoaiHoaDon.DIEN_NUOC.value,
                        HoaDon.ky_thanh_toan == thang,
                    )
                    .first()
                )
                candidates.append(
                    UtilityBillingCandidate(
                        ma_phong=item["ma_phong"],
                        so_phong=item["so_phong"],
                        toa_nha=item["toa_nha"],
                        so_sinh_vien=item["so_sinh_vien"],
                        chi_so_dien_cu_moi=item["chi_so_dien_cu_moi"],
                        so_dien_kwh=item["so_dien_kwh"],
                        chi_so_nuoc_cu_moi=item["chi_so_nuoc_cu_moi"],
                        so_nuoc_m3=item["so_nuoc_m3"],
                        tong_tien=item["tong_tien"],
                        da_lap_hoa_don=existing is not None,
                    )
                )

        return candidates

    @staticmethod
    def publish_utility_invoices(
        db: Session,
        request: UtilityInvoicePublishRequest,
        creator_username: str,
    ) -> PublishResultResponse:
        """
        Phát hành hóa đơn tiền điện nước theo tháng cho các phòng ký túc xá.
        Kiểm tra chống tạo trùng cho cùng phòng trong cùng tháng.
        """
        if not request.thang or not request.thang.strip():
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Tháng áp dụng không được để trống",
            )

        if not request.han_thanh_toan:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="Hạn chót thanh toán không được để trống",
            )

        all_candidates = InvoiceService.get_utility_billing_candidates(
            db=db,
            thang=request.thang,
        )

        targets = all_candidates
        if request.danh_sach_phong:
            selected_set = set(request.danh_sach_phong)
            targets = [c for c in all_candidates if c.ma_phong in selected_set]
            if not targets:
                raise HTTPException(
                    status_code=status.HTTP_404_NOT_FOUND,
                    detail="Không tìm thấy phòng nào phù hợp với danh sách mã phòng đã chỉ định",
                )

        created_invoices: List[HoaDon] = []
        already_issued_rooms: List[str] = []

        # Nếu có chi tiết phòng gửi kèm từ file chỉ số điện nước tải lên
        room_detail_map = {}
        if request.chi_tiet_phong:
            for item in request.chi_tiet_phong:
                room_detail_map[item.ma_phong.upper()] = item
                if item.so_phong:
                    room_detail_map[item.so_phong.upper()] = item

        for c in targets:
            existing = (
                db.query(HoaDon)
                .filter(
                    HoaDon.ma_phong == c.ma_phong,
                    HoaDon.loai_hoa_don == LoaiHoaDon.DIEN_NUOC.value,
                    HoaDon.ky_thanh_toan == request.thang,
                )
                .first()
            )

            if existing:
                already_issued_rooms.append(c.ma_phong)
                continue

            today_str = datetime.date.today().strftime("%Y%m%d")
            unique_suffix = uuid.uuid4().hex[:6].upper()
            ma_hd = f"HDDN-{today_str}-{unique_suffix}"

            # Tính toán tiền điện nước:
            meter_info = room_detail_map.get(c.ma_phong.upper()) or room_detail_map.get(c.so_phong.upper())
            if meter_info:
                calculated_amount = float(meter_info.tong_tien)
                so_dien_str = f"Điện: {meter_info.chi_so_dien_cu_moi} ({meter_info.so_dien_kwh} kWh)"
                so_nuoc_str = f"Nước: {meter_info.chi_so_nuoc_cu_moi} ({meter_info.so_nuoc_m3} m³)"
                note_str = f"Hóa đơn điện nước {request.thang} (Phòng {c.so_phong} - {c.toa_nha}. {so_dien_str}, {so_nuoc_str})"
            else:
                calculated_amount = (c.so_dien_kwh * request.don_gia_dien) + (c.so_nuoc_m3 * request.don_gia_nuoc)
                if calculated_amount <= 0:
                    calculated_amount = c.tong_tien
                note_str = request.ghi_chu or f"Hóa đơn tiền điện nước {request.thang} (Phòng {c.so_phong} - {c.toa_nha})"

            new_invoice = HoaDon(
                ma_hoa_don=ma_hd,
                ma_hop_dong=None,
                msv=None,
                ho_ten=None,
                ma_phong=c.ma_phong,
                so_phong=f"{c.so_phong} - {c.toa_nha}",
                loai_hoa_don=LoaiHoaDon.DIEN_NUOC.value,
                ky_thanh_toan=request.thang,
                so_tien=calculated_amount,
                ngay_lap=datetime.date.today(),
                han_thanh_toan=request.han_thanh_toan,
                trang_thai=TrangThaiHoaDon.CHUA_THANH_TOAN.value,
                ghi_chu=note_str,
                nguoi_tao=creator_username,
                ngay_tao=datetime.datetime.now(),
            )
            db.add(new_invoice)

            # Tạo bản ghi vào bảng so_cong_no
            so_cong_no_rec = SoCongNo(
                ma_hoa_don=new_invoice.ma_hoa_don,
                msv=None,
                so_phong=new_invoice.so_phong,
                loai_cong_no=new_invoice.loai_hoa_don,
                ky_thanh_toan=new_invoice.ky_thanh_toan,
                tong_tien=new_invoice.so_tien,
                da_tra=0.0,
                con_thieu=new_invoice.so_tien,
                han_thanh_toan=new_invoice.han_thanh_toan,
                trang_thai="CON_NO",
                ghi_chu=new_invoice.ghi_chu or f"Công nợ tiền điện nước {new_invoice.ky_thanh_toan}",
                ngay_cap_nhat=datetime.datetime.now(),
            )
            db.add(so_cong_no_rec)

            created_invoices.append(new_invoice)

        if not created_invoices and already_issued_rooms:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Tất cả các phòng được chọn đã được phát hành hóa đơn điện nước cho {request.thang} trước đó ({', '.join(already_issued_rooms)}). Không thể phát hành trùng!",
            )

        db.commit()

        for inv in created_invoices:
            db.refresh(inv)

        invoices_response = [
            HoaDonResponse(
                ma_hoa_don=inv.ma_hoa_don,
                ma_hop_dong=inv.ma_hop_dong,
                msv=inv.msv,
                ho_ten=inv.ho_ten,
                ma_phong=inv.ma_phong,
                so_phong=inv.so_phong,
                loai_hoa_don=inv.loai_hoa_don,
                ky_thanh_toan=inv.ky_thanh_toan,
                so_tien=inv.so_tien,
                ngay_lap=inv.ngay_lap,
                han_thanh_toan=inv.han_thanh_toan,
                trang_thai=inv.trang_thai,
                ghi_chu=inv.ghi_chu,
                nguoi_tao=inv.nguoi_tao,
            )
            for inv in created_invoices
        ]

        total_sum = sum(inv.so_tien for inv in created_invoices)
        msg = f"Đã phát hành thành công {len(created_invoices)} hóa đơn điện nước cho {request.thang}."
        if already_issued_rooms:
            msg += f" (Đã bỏ qua các phòng đã có hóa đơn: {', '.join(already_issued_rooms)})"

        return PublishResultResponse(
            success=True,
            message=msg,
            tong_hoa_don=len(created_invoices),
            tong_so_tien=total_sum,
            invoices=invoices_response,
        )

    @staticmethod
    def upload_utility_readings(
        db: Session,
        file_bytes: bytes,
        filename: str,
        thang: str = "Tháng 09/2026",
        don_gia_dien: float = 3000.0,
        don_gia_nuoc: float = 15000.0,
    ) -> UtilityMeterUploadResponse:
        """
        Đọc file Excel hoặc CSV chỉ số điện nước, tự động tính toán:
        - Số điện tiêu thụ (kWh) = Chỉ số mới - Chỉ số cũ
        - Số nước tiêu thụ (m3) = Chỉ số mới - Chỉ số cũ
        - Tiền điện = kWh * đơn giá điện
        - Tiền nước = m3 * đơn giá nước
        - Tổng tiền = Tiền điện + Tiền nước
        Trả về danh sách phòng kèm kết quả tính toán chi tiết.
        """
        import io, csv, re
        from openpyxl import load_workbook

        raw_rows = []
        is_excel = filename.lower().endswith((".xlsx", ".xls"))
        if is_excel:
            try:
                wb = load_workbook(io.BytesIO(file_bytes), data_only=True)
                ws = wb.active
                for row in ws.iter_rows(values_only=True):
                    if row and any(c is not None and str(c).strip() != "" for c in row):
                        raw_rows.append([str(c).strip() if c is not None else "" for c in row])
            except Exception as e:
                raise HTTPException(
                    status_code=status.HTTP_400_BAD_REQUEST,
                    detail=f"Không thể đọc file Excel chỉ số: {str(e)}",
                )
        else:
            decoded = None
            for enc in ["utf-8-sig", "utf-8", "cp1258", "latin1"]:
                try:
                    decoded = file_bytes.decode(enc)
                    break
                except UnicodeDecodeError:
                    continue
            if decoded is None:
                decoded = file_bytes.decode("utf-8", errors="ignore")
            delim = ";" if decoded.count(";") > decoded.count(",") else ","
            reader = csv.reader(io.StringIO(decoded), delimiter=delim)
            for row in reader:
                if row and any(c.strip() for c in row):
                    raw_rows.append([c.strip() for c in row])

        if not raw_rows:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="File chỉ số trống hoặc không có dòng dữ liệu hợp lệ",
            )

        header_idx = -1
        col_room = -1
        col_elec_old = -1
        col_elec_new = -1
        col_water_old = -1
        col_water_new = -1

        for idx, r in enumerate(raw_rows[:10]):
            r_lower = [str(c).lower().replace(" ", "").replace("_", "") for c in r]
            for c_idx, cell in enumerate(r_lower):
                if any(k in cell for k in ["phong", "maphong", "sophong", "room"]):
                    col_room = c_idx
                elif any(k in cell for k in ["diencu", "chisodiencu", "oldelec", "dien_cu"]):
                    col_elec_old = c_idx
                elif any(k in cell for k in ["dienmoi", "chisodienmoi", "newelec", "dien_moi"]):
                    col_elec_new = c_idx
                elif any(k in cell for k in ["nuoccu", "chisonuoccu", "oldwater", "nuoc_cu"]):
                    col_water_old = c_idx
                elif any(k in cell for k in ["nuocmoi", "chisonuocmoi", "newwater", "nuoc_moi"]):
                    col_water_new = c_idx

            if col_room != -1 and ((col_elec_old != -1 and col_elec_new != -1) or (col_water_old != -1 and col_water_new != -1)):
                header_idx = idx
                break

        if header_idx == -1:
            header_idx = 0
            if len(raw_rows[0]) >= 6:
                col_room = 1
                col_elec_old = 2
                col_elec_new = 3
                col_water_old = 4
                col_water_new = 5
            elif len(raw_rows[0]) >= 5:
                col_room = 0
                col_elec_old = 1
                col_elec_new = 2
                col_water_old = 3
                col_water_new = 4

        # Danh sách phòng trong DB để đối chiếu
        db_rooms = db.query(Phong).all()
        room_map = {}
        for r in db_rooms:
            toa = r.tang.toa_nha.ten_toa if r.tang and r.tang.toa_nha else "Tòa A1"
            room_map[r.ma_phong.upper()] = (r.ma_phong, r.so_phong, toa)
            room_map[r.so_phong.upper()] = (r.ma_phong, r.so_phong, toa)
            room_map[f"P{r.so_phong}".upper()] = (r.ma_phong, r.so_phong, toa)

        items: List[UtilityBillingCandidate] = []
        total_amount = 0.0

        for r_idx in range(header_idx + 1, len(raw_rows)):
            row = raw_rows[r_idx]
            if not row or col_room >= len(row) or not row[col_room].strip():
                continue

            raw_room = row[col_room].strip()
            if any(k in raw_room.lower() for k in ["phòng", "phong", "tổng", "stt"]):
                continue

            def parse_num(val_str: str) -> int:
                m = re.findall(r"\d+", val_str.replace(",", "").replace(".", ""))
                return int(m[0]) if m else 0

            elec_old = parse_num(row[col_elec_old]) if col_elec_old < len(row) and col_elec_old != -1 else 1000
            elec_new = parse_num(row[col_elec_new]) if col_elec_new < len(row) and col_elec_new != -1 else elec_old + 100
            water_old = parse_num(row[col_water_old]) if col_water_old < len(row) and col_water_old != -1 else 400
            water_new = parse_num(row[col_water_new]) if col_water_new < len(row) and col_water_new != -1 else water_old + 10

            elec_usage = max(0, elec_new - elec_old)
            water_usage = max(0, water_new - water_old)
            tien_dien = elec_usage * don_gia_dien
            tien_nuoc = water_usage * don_gia_nuoc
            row_total = tien_dien + tien_nuoc
            total_amount += row_total

            clean_r = re.sub(r"[^a-zA-Z0-9]", "", raw_room).upper()
            found = room_map.get(clean_r) or room_map.get(raw_room.upper())
            digits = re.findall(r"\d+", raw_room)
            if not found:
                if digits and digits[0] in room_map:
                    found = room_map[digits[0]]
                elif digits and f"P{digits[0]}" in room_map:
                    found = room_map[f"P{digits[0]}"]

            if found:
                ma_p, so_p, toa_p = found
            else:
                ma_p = clean_r or f"P{r_idx}"
                so_p = digits[0] if digits else raw_room
                toa_p = "Tòa A1"

            existing = (
                db.query(HoaDon)
                .filter(
                    HoaDon.ma_phong == ma_p,
                    HoaDon.loai_hoa_don == LoaiHoaDon.DIEN_NUOC.value,
                    HoaDon.ky_thanh_toan == thang,
                )
                .first()
            )

            items.append(
                UtilityBillingCandidate(
                    ma_phong=ma_p,
                    so_phong=so_p,
                    toa_nha=toa_p,
                    so_sinh_vien=4,
                    chi_so_dien_cu_moi=f"{elec_old} - {elec_new}",
                    so_dien_kwh=elec_usage,
                    chi_so_nuoc_cu_moi=f"{water_old} - {water_new}",
                    so_nuoc_m3=water_usage,
                    tong_tien=row_total,
                    so_dien_cu=elec_old,
                    so_dien_moi=elec_new,
                    so_nuoc_cu=water_old,
                    so_nuoc_moi=water_new,
                    tien_dien=tien_dien,
                    tien_nuoc=tien_nuoc,
                    da_lap_hoa_don=existing is not None,
                )
            )

        return UtilityMeterUploadResponse(
            success=True,
            message=f"Đã tính toán thành công chỉ số điện nước cho {len(items)} phòng từ file {filename}",
            fileName=filename,
            totalRooms=len(items),
            totalAmount=total_amount,
            items=items,
        )

    @staticmethod
    def generate_utility_template() -> bytes:
        """
        Tạo file Excel mẫu nhập chỉ số điện nước KTX
        """
        import openpyxl, io
        from openpyxl.styles import Font, Alignment, PatternFill

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Chỉ số điện nước"

        headers = [
            "STT",
            "Phòng",
            "Chỉ số điện cũ (kWh)",
            "Chỉ số điện mới (kWh)",
            "Chỉ số nước cũ (m³)",
            "Chỉ số nước mới (m³)",
            "Ghi chú",
        ]
        ws.append(headers)

        sample_data = [
            [1, "P101 - Tòa A1", 1000, 1100, 400, 410, "Tháng 09/2026"],
            [2, "P102 - Tòa A1", 1000, 1100, 400, 410, "Tháng 09/2026"],
            [3, "P203 - Tòa A1", 1000, 1100, 400, 410, "Tháng 09/2026"],
            [4, "P205 - Tòa A2", 1540, 1670, 540, 554, "Tháng 09/2026"],
        ]
        for row in sample_data:
            ws.append(row)

        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="2B85EC", end_color="2B85EC", fill_type="solid")
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.column_dimensions["A"].width = 8
        ws.column_dimensions["B"].width = 20
        ws.column_dimensions["C"].width = 24
        ws.column_dimensions["D"].width = 24
        ws.column_dimensions["E"].width = 24
        ws.column_dimensions["F"].width = 24
        ws.column_dimensions["G"].width = 18

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    @staticmethod
    def get_invoices(
        db: Session,
        loai_hoa_don: Optional[str] = None,
        ky_thanh_toan: Optional[str] = None,
        trang_thai: Optional[str] = None,
        msv: Optional[str] = None,
        ma_phong: Optional[str] = None,
    ) -> List[HoaDonResponse]:
        """
        Truy vấn danh sách hóa đơn theo các tiêu chí lọc.
        """
        query = db.query(HoaDon)

        if loai_hoa_don:
            query = query.filter(HoaDon.loai_hoa_don == loai_hoa_don)
        if ky_thanh_toan:
            query = query.filter(HoaDon.ky_thanh_toan == ky_thanh_toan)
        if trang_thai:
            query = query.filter(HoaDon.trang_thai == trang_thai)
        if msv:
            query = query.filter(HoaDon.msv == msv)
        if ma_phong:
            query = query.filter(HoaDon.ma_phong == ma_phong)

        invoices = query.order_by(HoaDon.ngay_tao.desc()).all()

        results: List[HoaDonResponse] = []
        for inv in invoices:
            results.append(
                HoaDonResponse(
                    ma_hoa_don=inv.ma_hoa_don,
                    ma_hop_dong=inv.ma_hop_dong,
                    msv=inv.msv,
                    ho_ten=inv.ho_ten,
                    ma_phong=inv.ma_phong,
                    so_phong=inv.so_phong,
                    loai_hoa_don=inv.loai_hoa_don,
                    ky_thanh_toan=inv.ky_thanh_toan,
                    so_tien=inv.so_tien,
                    ngay_lap=inv.ngay_lap,
                    han_thanh_toan=inv.han_thanh_toan,
                    trang_thai=inv.trang_thai,
                    ghi_chu=inv.ghi_chu,
                    nguoi_tao=inv.nguoi_tao,
                )
            )

        return results

    @staticmethod
    def get_invoice_by_id(
        db: Session,
        ma_hoa_don: str,
    ) -> Optional[HoaDonResponse]:
        """
        Lấy thông tin chi tiết một hóa đơn theo mã.
        """
        inv = db.query(HoaDon).filter(HoaDon.ma_hoa_don == ma_hoa_don).first()
        if not inv:
            return None

        return HoaDonResponse(
            ma_hoa_don=inv.ma_hoa_don,
            ma_hop_dong=inv.ma_hop_dong,
            msv=inv.msv,
            ho_ten=inv.ho_ten,
            ma_phong=inv.ma_phong,
            so_phong=inv.so_phong,
            loai_hoa_don=inv.loai_hoa_don,
            ky_thanh_toan=inv.ky_thanh_toan,
            so_tien=inv.so_tien,
            ngay_lap=inv.ngay_lap,
            han_thanh_toan=inv.han_thanh_toan,
            trang_thai=inv.trang_thai,
            ghi_chu=inv.ghi_chu,
            nguoi_tao=inv.nguoi_tao,
        )

